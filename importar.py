"""Importa o modelo original sem modificar a planilha de origem."""
from contextlib import closing
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

import pandas as pd
import database as db

DIAS = ('Segunda-feira', 'Terça-feira', 'Quarta-feira', 'Quinta-feira',
        'Sexta-feira', 'Sábado', 'Domingo')


def novo_resultado():
    return dict(operadores=0, produtos=0, programacoes=0, alocacoes=0,
                fora_maquina=0, ignorados=0, erros=[], avisos=[])


def _safe_str(val):
    if val is None or pd.isna(val):
        return None
    s = str(val).strip()
    return s if s and s.lower() not in ('nan', 'none', 'nat') and not s.startswith('#') else None


def _safe_float_str(val):
    """Mantém zeros textuais e não trunca frações de identificadores."""
    s = _safe_str(val)
    if not s:
        return None
    if re.fullmatch(r'\d+', s):
        return s
    try:
        n = Decimal(s.replace(',', '.'))
        return str(int(n)) if n.is_finite() and n >= 0 and n == n.to_integral() else None
    except (ValueError, InvalidOperation):
        return None


def _codigo(val):
    s = _safe_float_str(val) if isinstance(val, (int, float)) else _safe_str(val)
    return s.upper() if s else None


def _normalizar(val):
    return ' '.join(unicodedata.normalize('NFKD', str(val)).encode(
        'ascii', 'ignore').decode().lower().split())


def _celula(df, linha, coluna):
    return df.iat[linha, coluna] if linha < len(df) and coluna < len(df.columns) else None


def _data(val):
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (date, datetime)):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, (int, float)):
        if not 1 <= val <= 100000:
            return None
        return (datetime(1899, 12, 30) + timedelta(days=float(val))).strftime('%Y-%m-%d')
    for formato in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d-%m-%Y', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(str(val).strip(), formato).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def _ler_aba(xl, aba, mesclas):
    df = pd.read_excel(xl, sheet_name=aba, header=None, dtype=object)
    df.attrs['linhas_com_dados'] = {
        lado: {i for i in range(len(df)) if any(_safe_str(_celula(df, i, c))
               for c in (lado+1, lado+3, lado+5))} for lado in (1, 8)
    }
    ranges = mesclas.get(aba, []) if xl.engine == 'openpyxl' else xl.book.sheet_by_name(aba).merged_cells
    # Replica somente células realmente mescladas de máquina/produto/OP.
    for r0, r1, c0, c1 in ranges:
        if c1-c0 == 1 and c0 in (1, 4, 5, 6, 8, 11, 12, 13):
            val = _celula(df, r0, c0)
            if val is not None and c0 < len(df.columns):
                df.iloc[r0:min(r1, len(df)), c0] = val
    return df


def _cadastro(conn, tabela, chave, codigo, campo, texto, resultados):
    # Identificadores SQL definidos exclusivamente pelas chamadas internas.
    rows = conn.execute(f'SELECT * FROM {tabela} WHERE {chave}=? COLLATE NOCASE ORDER BY ({chave}=?) DESC,id', (codigo, codigo)).fetchall()
    if len(rows) > 1:
        aviso = f'{codigo}: existem cadastros antigos com maiúsculas/minúsculas diferentes. Usado cadastro {rows[0]["id"]}; revise as duplicatas.'
        if aviso not in resultados['avisos']:
            resultados['avisos'].append(aviso)
    if rows:
        return rows[0]['id']
    if not texto:
        raise ValueError(f'{codigo}: {campo} ausente e código não encontrado no cadastro.')
    cur = conn.execute(f'INSERT INTO {tabela} ({chave},{campo}) VALUES (?,?)', (codigo, texto))
    resultados[tabela] += 1
    return cur.lastrowid


def _operador(conn, cracha, nome, resultados):
    return _cadastro(conn, 'operadores', 'cracha', cracha, 'nome', nome.upper() if nome else None,
                     resultados) if cracha else None


def _produto(conn, codigo, descricao, resultados):
    return _cadastro(conn, 'produtos', 'codigo', codigo, 'descricao', descricao,
                     resultados) if codigo else None


def _linha(conn, df, i, lado, prog_id, resultados):
    if i not in df.attrs['linhas_com_dados'][lado]:
        return
    codigo = _codigo(_celula(df, i, lado))
    if not codigo or not re.fullmatch(r'(?:PCT|LXR|LXD|SCL|CSD)\s*\d+', codigo):
        return
    codigo = codigo.replace(' ', '')
    cracha_val = _celula(df, i, lado+1)
    cracha = _safe_float_str(cracha_val)
    nome = _safe_str(_celula(df, i, lado+2))
    produto = _codigo(_celula(df, i, lado+3))
    descricao = _safe_str(_celula(df, i, lado+4))
    op_val = _celula(df, i, lado+5)
    num_op = _safe_float_str(op_val)
    if _safe_str(cracha_val) and not cracha:
        raise ValueError(f'Crachá inválido: {cracha_val}')
    if _safe_str(op_val) and not num_op:
        raise ValueError(f'Número de OP inválido: {op_val}')
    if not (cracha or produto or num_op):
        return
    op_id = _operador(conn, cracha, nome, resultados)
    prod_id = _produto(conn, produto, descricao, resultados)
    conn.execute('INSERT OR IGNORE INTO maquinas (codigo,setor,descricao) VALUES (?,?,?)',
                 (codigo, _detectar_setor(codigo), ''))
    maq_id = conn.execute('SELECT id FROM maquinas WHERE codigo=?', (codigo,)).fetchone()[0]
    valores = (prog_id, maq_id, op_id, prod_id, num_op or '')
    existe = conn.execute("""SELECT a.id FROM alocacoes a LEFT JOIN produtos p ON p.id=a.produto_id
        WHERE a.programacao_id=? AND a.maquina_id=? AND a.operador_id IS ?
        AND (a.produto_id IS ? OR p.codigo=? COLLATE NOCASE) AND COALESCE(a.num_op,'')=?""",
        (prog_id, maq_id, op_id, prod_id, produto, num_op or '')).fetchone()
    if existe:
        resultados['ignorados'] += 1
    else:
        conn.execute('''INSERT INTO alocacoes
            (programacao_id,maquina_id,operador_id,produto_id,num_op) VALUES (?,?,?,?,?)''', valores)
        resultados['alocacoes'] += 1


def _fora(conn, df, i, prog_id, resultados):
    val = _celula(df, i, 2)
    # O modelo também contém crachá digitado em B em vez de C.
    if not _safe_str(val):
        val = _celula(df, i, 1)
    cracha = _safe_float_str(val)
    nome = _safe_str(_celula(df, i, 3))
    obs = _safe_str(_celula(df, i, 4)) or ''
    if not cracha:
        if _safe_str(val) or nome or obs:
            raise ValueError('Fora de máquina: informe um crachá válido.')
        return
    op_id = _operador(conn, cracha, nome, resultados)
    valores = (prog_id, op_id, obs)
    if conn.execute("""SELECT id FROM fora_maquina WHERE programacao_id=?
        AND operador_id=? AND COALESCE(observacao,'')=?""", valores).fetchone():
        resultados['ignorados'] += 1
    else:
        conn.execute('INSERT INTO fora_maquina (programacao_id,operador_id,observacao) VALUES (?,?,?)', valores)
        resultados['fora_maquina'] += 1


def _tentar(conn, resultados, contexto, acao):
    contagens = {k: v for k, v in resultados.items() if isinstance(v, int)}
    conn.execute('SAVEPOINT linha_importada')
    try:
        acao()
    except (ValueError, db.sqlite3.IntegrityError) as exc:
        conn.execute('ROLLBACK TO linha_importada')
        resultados.update(contagens)
        resultados['erros'].append(f'{contexto}: {exc}')
    finally:
        conn.execute('RELEASE linha_importada')


def importar_planilha(caminho):
    resultados = novo_resultado()
    try:
        mesclas = {}
        if Path(caminho).suffix.lower() in ('.xlsx', '.xlsm'):
            from openpyxl import load_workbook
            with open(caminho, 'rb') as arquivo:
                workbook = load_workbook(arquivo, data_only=True)
                try:
                    mesclas = {s.title: [(r.min_row-1, r.max_row, r.min_col-1, r.max_col)
                                         for r in s.merged_cells.ranges] for s in workbook}
                finally:
                    workbook.close()
        kwargs = {'engine_kwargs': {'formatting_info': True}} if Path(caminho).suffix.lower() == '.xls' else {}
        with pd.ExcelFile(caminho, **kwargs) as xl, closing(db.get_conn()) as conn, conn:
            conn.execute('BEGIN IMMEDIATE')
            cadastros = [a for a in xl.sheet_names if _normalizar(a) in ('plan2 (2)', 'plan2')]
            turnos = [a for a in xl.sheet_names if re.match(r'^[123][o°º]?\s*turno\b', _normalizar(a))
                      or _normalizar(a) == 'domingo']
            if not cadastros and not turnos:
                raise ValueError('Nenhuma aba reconhecida. Esperado: Plan2 (2), 1º/2º/3º Turno ou Domingo.')
            for aba in cadastros:
                df = pd.read_excel(xl, sheet_name=aba, header=None, dtype=object)
                if len(df.columns) < 7:
                    raise ValueError(f'{aba}: esperado cadastro de operadores em C:D e produtos em F:G.')
                for i in range(len(df)):
                    cracha, nome = _safe_float_str(_celula(df, i, 2)), _safe_str(_celula(df, i, 3))
                    codigo, descricao = _codigo(_celula(df, i, 5)), _safe_str(_celula(df, i, 6))
                    if cracha and nome:
                        _tentar(conn, resultados, f'{aba}, linha {i+1}', lambda: _operador(conn, cracha, nome, resultados))
                    if codigo and descricao and _normalizar(codigo) != 'produto':
                        _tentar(conn, resultados, f'{aba}, linha {i+1}', lambda: _produto(conn, codigo, descricao, resultados))
            for aba in turnos:
                df = _ler_aba(xl, aba, mesclas)
                if not any(_normalizar(_celula(df, 4, c)) == 'maquina' for c in (1, 8)):
                    resultados['erros'].append(f'{aba}: cabeçalho Máquina não encontrado em B5/I5. Aba não importada.')
                    continue
                data_str = _data(_celula(df, 1, 12))
                if not data_str:
                    resultados['erros'].append(f'{aba}: data ausente/inválida em M2. Aba não importada.')
                    continue
                titulo_turno = re.search(r'turno\s*0?([123])\b', _normalizar(_celula(df, 1, 13)))
                turno = int(aba.strip()[0]) if aba.strip()[0] in '123' else (int(titulo_turno[1]) if titulo_turno else None)
                if not turno:
                    resultados['erros'].append(f'{aba}: turno não identificado em N2. Aba não importada.')
                    continue
                if titulo_turno and int(titulo_turno[1]) != turno:
                    resultados['avisos'].append(f'{aba}: N2 informa outro turno. Usado turno {turno} do nome da aba.')
                dia = DIAS[date.fromisoformat(data_str).weekday()]
                if ('sabado' in _normalizar(aba) and dia != 'Sábado') or (_normalizar(aba) == 'domingo' and dia != 'Domingo'):
                    resultados['avisos'].append(f'{aba}: M2 contém {data_str} ({dia}). Mantida a data da célula.')
                progs = conn.execute('SELECT id FROM programacoes WHERE data=? AND turno=?', (data_str, turno)).fetchall()
                if len(progs) > 1:
                    resultados['erros'].append(f'{aba}: programações duplicadas para {data_str}, turno {turno}. Revise antes de importar.')
                    continue
                if progs:
                    prog_id = progs[0][0]
                else:
                    prog_id = conn.execute('INSERT INTO programacoes (data,turno,dia_semana) VALUES (?,?,?)',
                                           (data_str, turno, dia)).lastrowid
                    resultados['programacoes'] += 1
                fora = False
                for i in range(4, len(df)):
                    if 'fora de maquina' in _normalizar(_celula(df, i, 2)):
                        fora = True
                    else:
                        _tentar(conn, resultados, f'{aba}, linha {i+1}, esquerda',
                                lambda: _fora(conn, df, i, prog_id, resultados) if fora else _linha(conn, df, i, 1, prog_id, resultados))
                    _tentar(conn, resultados, f'{aba}, linha {i+1}, direita',
                            lambda: _linha(conn, df, i, 8, prog_id, resultados))
    except Exception as exc:
        resultados = novo_resultado()
        resultados['erros'].append(f'Importação cancelada; nenhuma alteração gravada: {exc}')
    return resultados


def _detectar_setor(codigo):
    codigo = codigo.upper()
    if codigo.startswith('PCT'): return 'Picote'
    if codigo.startswith(('LXR', 'LXD')): return 'Saco para Lixo'
    if codigo.startswith('SCL'): return 'Sacoleira'
    if codigo.startswith('CSD'): return 'Corte Solda'
    return 'Outro'
