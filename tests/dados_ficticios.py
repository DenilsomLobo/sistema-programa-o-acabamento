"""Planilha de teste criada do zero; não lê arquivos ou cadastros da empresa."""
from datetime import datetime
from openpyxl import Workbook


def criar_planilha(caminho):
    wb = Workbook()
    cadastro = wb.active
    cadastro.title = 'Plan2 (2)'
    cadastro.append([None, None, 'Crachá', 'Nome', None, 'Produto', 'Descrição'])
    for i in range(1, 4):
        cadastro.append([None, None, str(900000+i), 'John Doe', None,
                         f'DEMO-{i:03d}', f'Produto fictício {i}'])
    for grupo, data in (('V2', datetime(2024, 1, 5)), ('sábado', datetime(2024, 1, 6))):
        for turno in (1, 2, 3):
            _turno(wb.create_sheet(f'{turno}º Turno {grupo}'), data, turno)
    # Data propositalmente incompatível com o título para testar o aviso.
    _turno(wb.create_sheet('Domingo'), datetime(2024, 1, 10), 3)
    wb.save(caminho)
    wb.close()
    return caminho


def _turno(s, data, turno):
    s['M2'] = data
    s['N2'] = f'Turno {turno:02d}'
    for col in (2, 9):
        for offset, titulo in enumerate(('Máquina', 'Crachá', 'Nome', 'Produto', 'Descrição', 'Nº da OP')):
            s.cell(5, col+offset, titulo)
        codigo = 'PCT900001' if col == 2 else 'LXR900002'
        valores = (codigo, '900001', 'John Doe', 'demo-001', 'Produto fictício 1', '800001')
        for offset, valor in enumerate(valores):
            s.cell(6, col+offset, valor)
        s.cell(7, col+1, '900002')
        s.cell(7, col+2, 'John Doe')
        for offset in (0, 3, 4, 5):
            s.merge_cells(start_row=6, end_row=8, start_column=col+offset, end_column=col+offset)
        # A linha 8 vazia não deve virar alocação adicional ao expandir mesclas.
    for col, valor in enumerate(('PCT900005', '900001', 'John Doe', 'DEMO-AUSENTE', '#N/A', '800002'), 2):
        s.cell(9, col, valor)
    s['C10'] = 'PROGRAMAÇÃO FORA DE MÁQUINA OU EM OUTRO SETOR'
    s['C11'] = '900003'
    s['D11'] = 'John Doe'
    s['E11'] = 'Atividade fictícia'
